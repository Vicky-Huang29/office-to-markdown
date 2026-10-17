#!/usr/bin/env python3
"""Convert DOCX/PPTX manuals to structured Markdown with local image assets."""

from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path
from zipfile import ZipFile

from lxml import etree


NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "v": "urn:schemas-microsoft-com:vml",
}
W = NS["w"]

CIRCLED = {
    1: "①",
    2: "②",
    3: "③",
    4: "④",
    5: "⑤",
    6: "⑥",
    7: "⑦",
    8: "⑧",
    9: "⑨",
    10: "⑩",
    11: "⑪",
    12: "⑫",
    13: "⑬",
    14: "⑭",
    15: "⑮",
    16: "⑯",
    17: "⑰",
    18: "⑱",
    19: "⑲",
    20: "⑳",
}


def first_val(el, xpath, default=None, ns=NS):
    vals = el.xpath(xpath, namespaces=ns)
    return vals[0] if vals else default


def collapse_blank_lines(lines: list[str]) -> list[str]:
    clean = []
    prev_blank = False
    for line in lines:
        blank = line.strip() == ""
        if blank and prev_blank:
            continue
        clean.append("" if blank else line)
        prev_blank = blank
    while clean and clean[0] == "":
        clean.pop(0)
    return clean


def clean_text(text: str) -> str:
    text = (text or "").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def text_lines(text: str) -> list[str]:
    return [line.strip() for line in clean_text(text).split("\n") if line.strip()]


def prepare_output(src: Path, out_dir: Path) -> tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    md_path = out_dir / f"{src.stem}.md"
    asset_dir = out_dir / f"{src.stem}_assets"
    if asset_dir.exists():
        shutil.rmtree(asset_dir)
    asset_dir.mkdir(parents=True, exist_ok=True)
    return md_path, asset_dir


def convert_docx(src: Path, out_dir: Path) -> tuple[Path, int]:
    md_path, asset_dir = prepare_output(src, out_dir)

    with ZipFile(src) as z:
        document = etree.fromstring(z.read("word/document.xml"))
        styles_xml = etree.fromstring(z.read("word/styles.xml"))
        rels_xml = etree.fromstring(z.read("word/_rels/document.xml.rels"))
        numbering = (
            etree.fromstring(z.read("word/numbering.xml"))
            if "word/numbering.xml" in z.namelist()
            else etree.Element("numbering")
        )
        zip_files = {name: z.read(name) for name in z.namelist() if name.startswith("word/media/")}

    rels_ns = {"pr": "http://schemas.openxmlformats.org/package/2006/relationships"}
    rel_targets = {rel.get("Id"): rel.get("Target") for rel in rels_xml.xpath("//pr:Relationship", namespaces=rels_ns)}

    style_names = {}
    for style in styles_xml.xpath("//w:style", namespaces=NS):
        sid = style.get(f"{{{W}}}styleId")
        name = first_val(style, "./w:name/@w:val")
        if sid and name:
            style_names[sid] = name

    num_to_abs = {}
    level_defs = {}
    for num in numbering.xpath("//w:num", namespaces=NS):
        num_id = num.get(f"{{{W}}}numId")
        abs_id = first_val(num, "./w:abstractNumId/@w:val")
        if num_id and abs_id:
            num_to_abs[num_id] = abs_id
    for absn in numbering.xpath("//w:abstractNum", namespaces=NS):
        abs_id = absn.get(f"{{{W}}}abstractNumId")
        levels = {}
        for lvl in absn.xpath("./w:lvl", namespaces=NS):
            ilvl = int(lvl.get(f"{{{W}}}ilvl") or 0)
            levels[ilvl] = {
                "start": int(first_val(lvl, "./w:start/@w:val", "1")),
                "fmt": first_val(lvl, "./w:numFmt/@w:val", "decimal"),
                "text": first_val(lvl, "./w:lvlText/@w:val", "%1"),
            }
        if abs_id:
            level_defs[abs_id] = levels

    counters = {}
    image_index = 0
    lines: list[str] = []

    def format_number(n: int, fmt: str) -> str:
        if fmt == "lowerLetter":
            s = ""
            x = n
            while x:
                x -= 1
                s = chr(ord("a") + x % 26) + s
                x //= 26
            return s or "a"
        if fmt == "upperLetter":
            return format_number(n, "lowerLetter").upper()
        if fmt == "decimalEnclosedCircleChinese":
            return CIRCLED.get(n, f"({n})")
        return str(n)

    def next_label(num_id: str | None, ilvl: int) -> str:
        if not num_id or num_id == "0" or num_id not in num_to_abs:
            return ""
        abs_id = num_to_abs[num_id]
        levels = level_defs.get(abs_id, {})
        if ilvl not in levels:
            return ""
        state = counters.setdefault(num_id, {})
        for lvl in range(ilvl):
            if lvl not in state and lvl in levels:
                state[lvl] = levels[lvl]["start"]
        start = levels[ilvl]["start"]
        state[ilvl] = state.get(ilvl, start - 1) + 1
        for lvl in list(state):
            if lvl > ilvl:
                del state[lvl]

        def repl(match):
            lvl = int(match.group(1)) - 1
            val = state.get(lvl, levels.get(lvl, {"start": 1})["start"])
            fmt = levels.get(lvl, levels[ilvl])["fmt"]
            return format_number(val, fmt)

        label = re.sub(r"%([1-9])", repl, levels[ilvl]["text"])
        return label.replace("(", "（").replace(")", "）").strip()

    def add_line(line: str = ""):
        if lines and lines[-1] != "":
            lines.append("")
        if line:
            lines.append(line)

    def paragraph_style(p):
        sid = first_val(p, "./w:pPr/w:pStyle/@w:val")
        return style_names.get(sid, sid or "")

    def paragraph_num(p):
        num_id = first_val(p, "./w:pPr/w:numPr/w:numId/@w:val")
        ilvl = first_val(p, "./w:pPr/w:numPr/w:ilvl/@w:val")
        return num_id, int(ilvl or 0)

    def text_from(node):
        parts = []
        for el in node.iter():
            if el.tag == f"{{{W}}}t":
                parts.append(el.text or "")
            elif el.tag == f"{{{W}}}tab":
                parts.append("\t")
            elif el.tag == f"{{{W}}}br":
                parts.append("\n")
        return "".join(parts).strip()

    def md_cell(text: str) -> str:
        text = clean_text(text)
        text = re.sub(r"\s*\n\s*", "<br>", text)
        return text.replace("|", "\\|")

    def table_rows(tbl) -> list[list[str]]:
        rows = []
        for tr in tbl.xpath("./w:tr", namespaces=NS):
            row = []
            for tc in tr.xpath("./w:tc", namespaces=NS):
                parts = []
                for p in tc.xpath("./w:p", namespaces=NS):
                    value = text_from(p)
                    if value:
                        parts.append(value)
                row.append("\n".join(parts))
            rows.append(row)
        while rows and all(not cell.strip() for cell in rows[-1]):
            rows.pop()
        if not rows:
            return []
        width = max((len(row) for row in rows), default=0)
        return [row + [""] * (width - len(row)) for row in rows]

    def add_table(tbl):
        rows = table_rows(tbl)
        if not rows:
            return
        width = max(len(row) for row in rows)
        header = rows[0]
        body = rows[1:]
        if all(not cell.strip() for cell in header):
            header = [f"列{i}" for i in range(1, width + 1)]
            body = rows
        if lines and lines[-1] != "":
            lines.append("")
        lines.append("| " + " | ".join(md_cell(cell) or " " for cell in header) + " |")
        lines.append("| " + " | ".join("---" for _ in range(width)) + " |")
        for row in body:
            lines.append("| " + " | ".join(md_cell(cell) for cell in row) + " |")

    def save_image(rid: str) -> str | None:
        nonlocal image_index
        target = rel_targets.get(rid)
        if not target:
            return None
        target = target.replace("\\", "/")
        if target.startswith("../"):
            name = target[3:]
        elif target.startswith("word/"):
            name = target
        else:
            name = "word/" + target
        data = zip_files[name]
        ext = Path(name).suffix or ".bin"
        image_index += 1
        filename = f"image-{image_index:03d}{ext.lower()}"
        (asset_dir / filename).write_bytes(data)
        return f"{asset_dir.name}/{filename}"

    body = document.find("w:body", namespaces=NS)
    for child in body:
        if child.tag == f"{{{W}}}p":
            text = text_from(child)
            refs = child.xpath(".//a:blip/@r:embed", namespaces=NS) + child.xpath(".//v:imagedata/@r:id", namespaces=NS)
            style = paragraph_style(child)
            num_id, ilvl = paragraph_num(child)
            label = next_label(num_id, ilvl) if text else ""
            heading_match = re.match(r"heading\s+(\d+)", style or "", re.I)
            if text:
                if heading_match:
                    level = min(int(heading_match.group(1)), 6)
                    add_line(f"{'#' * level} {f'{label} {text}'.strip() if label else text}")
                elif label:
                    add_line(f"{'  ' * ilvl}{label} {text}")
                else:
                    add_line(text)
            for rid in refs:
                path = save_image(rid)
                if path:
                    add_line(f"![]({path})")
        elif child.tag == f"{{{W}}}tbl":
            add_table(child)

    md_path.write_text("\n".join(collapse_blank_lines(lines)).rstrip() + "\n", encoding="utf-8")
    return md_path, image_index


def convert_pptx(src: Path, out_dir: Path, skip_transition_slides: bool = False) -> tuple[Path, int]:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise SystemExit("Missing dependency: python -m pip install python-pptx") from exc

    md_path, asset_dir = prepare_output(src, out_dir)
    prs = Presentation(str(src))
    image_index = 0
    lines: list[str] = []

    def add_line(line: str = ""):
        if lines and lines[-1] != "":
            lines.append("")
        if line:
            lines.append(line)

    def paragraph_prefix(paragraph, counters):
        ppr = paragraph._p.pPr
        if ppr is None:
            return ""
        auto = ppr.find(f"{{{NS['a']}}}buAutoNum")
        if auto is None:
            return ""
        auto_type = auto.get("type")
        counters[auto_type] = counters.get(auto_type, 0) + 1
        n = counters[auto_type]
        if auto_type == "circleNumDbPlain":
            return CIRCLED.get(n, f"{n}.") + " "
        return f"{n}. "

    def text_from_shape(shape):
        if not getattr(shape, "has_text_frame", False):
            return ""
        counters = {}
        chunks = []
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if not text:
                continue
            prefix = paragraph_prefix(paragraph, counters)
            level = getattr(paragraph, "level", 0) or 0
            chunks.append(("  " * level) + prefix + text)
        return clean_text("\n".join(chunks))

    def text_from_shapes(shapes):
        parts = []
        for shape in sorted(shapes, key=lambda s: ((getattr(s, "top", 0) or 0), (getattr(s, "left", 0) or 0))):
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                nested = text_from_shapes(shape.shapes)
                if nested:
                    parts.append(nested)
            else:
                text = text_from_shape(shape)
                if text:
                    parts.append(text)
        return clean_text("\n".join(parts))

    def save_picture(shape):
        nonlocal image_index
        image_index += 1
        filename = f"image-{image_index:03d}.{(shape.image.ext or 'png').lower()}"
        (asset_dir / filename).write_bytes(shape.image.blob)
        return f"{asset_dir.name}/{filename}"

    def collect_items(shapes, skip_elements=None):
        skip_elements = skip_elements or set()
        items = []
        for idx, shape in enumerate(shapes):
            if getattr(shape, "_element", None) in skip_elements:
                continue
            top = getattr(shape, "top", 0) or 0
            left = getattr(shape, "left", 0) or 0
            width = getattr(shape, "width", 0) or 0
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                text = text_from_shapes(shape.shapes)
                if text:
                    items.append(("text", top, left, width, idx, text))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                items.append(("image", top, left, width, idx, save_picture(shape)))
            elif getattr(shape, "has_text_frame", False):
                text = text_from_shape(shape)
                if text:
                    items.append(("text", top, left, width, idx, text))
        return items

    def render_item(item):
        if item[0] == "text":
            for line in text_lines(item[5]):
                add_line(line)
        elif item[0] == "image":
            add_line(f"![]({item[5]})")

    def is_transition_slide(slide, title):
        texts = []
        pics = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pics += 1
            text = text_from_shape(shape)
            if text and text != title:
                texts.extend(text_lines(text))
        if pics:
            return False
        meaningful = [t for t in texts if t.strip()]
        return bool(meaningful and len(meaningful) <= 2 and any(re.fullmatch(r"\d+", t) for t in meaningful))

    def render_label_value(items, known_labels=None):
        known_labels = set(known_labels or [])
        labels = []
        values = []
        for item in items:
            if item[0] != "text":
                continue
            text = clean_text(item[5])
            if known_labels:
                if text in known_labels:
                    labels.append(item)
                elif text not in {"业务要点"}:
                    values.append(item)
            elif "\n" not in text and len(text) <= 12 and not re.match(r"^\d", text):
                labels.append(item)
            else:
                values.append(item)
        used = set()
        for label in sorted(labels, key=lambda x: (x[1], x[2])):
            candidates = []
            for value in values:
                if id(value) in used or value[1] <= label[1]:
                    continue
                if value[1] - label[1] > 1400000:
                    continue
                lx = label[2] + label[3] / 2
                vx = value[2] + value[3] / 2
                candidates.append((value[1] - label[1] + abs(vx - lx) * 0.35, value))
            if candidates:
                value = min(candidates, key=lambda x: x[0])[1]
                used.add(id(value))
                add_line(f"{clean_text(label[5])}：{'；'.join(text_lines(value[5]))}")
            else:
                render_item(label)
        for item in sorted(items, key=lambda x: (x[1], x[2])):
            if item[0] == "text" and clean_text(item[5]) in known_labels.union({"业务要点"}):
                continue
            if item[0] != "text" or (id(item) not in used and item not in labels):
                render_item(item)

    for slide_index, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = clean_text(text_from_shape(title_shape)) if title_shape is not None else ""
        if skip_transition_slides and is_transition_slide(slide, title):
            continue
        add_line(f"# 幻灯片 {slide_index}" + (f"：{title}" if title else ""))
        skip = {title_shape._element} if title_shape is not None else set()
        items = collect_items(slide.shapes, skip)
        if title == "业务要点":
            render_label_value(
                items,
                {
                    "票据类型",
                    "账户要求",
                    "提票渠道",
                    "融资要求",
                    "持票人属地要求",
                    "营业时间",
                },
            )
        else:
            for item in sorted(items, key=lambda x: (x[1], x[2])):
                render_item(item)

    md_path.write_text("\n".join(collapse_blank_lines(lines)).rstrip() + "\n", encoding="utf-8")
    return md_path, image_index


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--skip-transition-slides", action="store_true")
    args = parser.parse_args()

    suffix = args.input.suffix.lower()
    if suffix == ".docx":
        md_path, images = convert_docx(args.input, args.output_dir)
    elif suffix == ".pptx":
        md_path, images = convert_pptx(args.input, args.output_dir, args.skip_transition_slides)
    else:
        raise SystemExit(f"Unsupported file type: {suffix}")

    print(md_path)
    print(f"images={images}")


if __name__ == "__main__":
    main()
